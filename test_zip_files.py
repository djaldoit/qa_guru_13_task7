import zipfile
from pypdf import PdfReader
from openpyxl import load_workbook
from conftest import zip_resources


def test_pdf_file():
    with zipfile.ZipFile(zip_resources) as file_zip:
        with file_zip.open('file_pdf.pdf') as text_pdf:
            reader = PdfReader(text_pdf)
            text = reader.pages[0].extract_text()
            assert 'Y ou will need Adobe Acrobat Reader.' in text
            print('\nПроверка текста pdf-файла прошла успешно!')


def test_csv_file():
    with zipfile.ZipFile(zip_resources) as zip_file:
        with zip_file.open('file_csv.csv') as csv_file:
            reader = csv_file.read().decode('utf-8-sig')
            assert 'Elena' in reader



def test_xlsx_file():
    with zipfile.ZipFile(zip_resources) as zip_file:
        with zip_file.open('file_xlsx.xlsx') as file:
            workbook = load_workbook(file)
            sheet = workbook.active
            assert sheet.cell(row=1, column=1).value == '1 строка 1 столбец'
            assert sheet['A2'].value == '2 строка 1 столбец'